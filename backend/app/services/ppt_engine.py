import os
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.config import settings


def analyze_ppt_presentation(ppt_path: Optional[str], submission_id: Optional[str] = None) -> Dict[str, Any]:
    """Module 10: PPT Analysis - real per-slide text extraction via
    python-pptx. Deterministic extraction only; qualitative judgment
    (missing sections, presentation quality, improvement suggestions) moves
    into the Gemini prompt in Phase 7 rather than being faked locally here.

    `submission_id`, if given, also saves actual image bytes for any
    picture shapes (not just a has_image flag) to the shared workspace -
    used by the Pitch-Deck Reasoning Agent (Phase 14) to multimodally
    inspect slides whose extracted text alone can't establish their role in
    the pitch (an architecture diagram, a screenshot, a demo frame)."""
    if not ppt_path or not os.path.exists(ppt_path):
        return {"status": "skipped", "reason": "No presentation provided", "slide_count": 0, "slides": []}

    try:
        presentation = Presentation(ppt_path)
    except Exception as e:
        return {"status": "error", "reason": f"Could not open presentation: {e}", "slide_count": 0, "slides": []}

    slides_dir = None
    if submission_id:
        slides_dir = os.path.join(settings.WORKSPACE_ROOT, submission_id, "slides")
        os.makedirs(slides_dir, exist_ok=True)

    slides: List[Dict[str, Any]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = []
        has_image = False
        image_paths: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    texts.append(text)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                has_image = True
                if slides_dir:
                    try:
                        ext = shape.image.ext or "png"
                        image_path = os.path.join(slides_dir, f"slide_{index}_{shape.shape_id}.{ext}")
                        with open(image_path, "wb") as f:
                            f.write(shape.image.blob)
                        image_paths.append(image_path)
                    except (AttributeError, OSError):
                        pass  # some picture shapes (e.g. placeholders) have no readable image blob
        slides.append({
            "slide_number": index,
            "text": "\n".join(texts),
            "has_image": has_image,
            "image_paths": image_paths,
        })

    return {
        "status": "completed",
        "slide_count": len(slides),
        "slides": slides,
    }
